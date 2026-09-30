#!/usr/bin/env python3
"""複数の学生のログをまとめて一覧にする。

IoU 0.75 を主指標、0.5 を参考値として並べて出す。
ヘッダに MODEL / INPUT / FPS があればそれも表示する。
古いログ（ヘッダに新項目がないもの）はファイル名から推定するか空欄にする。

使い方:
  python report.py log1.txt log2.txt log3.txt
  python report.py logs/*.txt
  python report.py logs/*.txt --tsv report.tsv    # TSV でも出力
  python report.py logs/*.txt --xlsx report.xlsx  # Excel でも出力
"""

import argparse
import os
import re
import sys
import unicodedata

import yaml

# eval_log.py の関数を再利用
sys.path.insert(0, os.path.dirname(__file__))
from eval_log import (
    load_config, parse_log, check_header,
    build_frame_results, build_held_results,
    write_detection_results, run_map,
    auto_scale_if_needed,
)


# ---------- ユーティリティ ----------

def east_asian_width(s):
    """文字列の端末表示幅を返す（全角=2, 半角=1）。"""
    w = 0
    for c in s:
        eaw = unicodedata.east_asian_width(c)
        w += 2 if eaw in ("F", "W", "A") else 1
    return w


def pad_to_width(s, width, align="right"):
    """端末表示幅が width になるよう空白で埋める。"""
    cur = east_asian_width(s)
    pad = max(0, width - cur)
    if align == "left":
        return s + " " * pad
    return " " * pad + s


def shorten_model(name, max_width=22):
    """モデル名を max_width 以内に収める。.xmodel を省く。"""
    if not name:
        return ""
    # .xmodel を除去
    name = re.sub(r'\.xmodel$', '', name)
    if east_asian_width(name) <= max_width:
        return name
    # それでも長ければ末尾を切る
    result = ""
    for c in name:
        eaw = unicodedata.east_asian_width(c)
        cw = 2 if eaw in ("F", "W", "A") else 1
        if east_asian_width(result) + cw + 1 > max_width:
            return result + "~"
        result += c
    return result


# ---------- モデル・入力サイズの推定 ----------

def guess_model_from_filename(filename):
    """ファイル名からモデル名を推定する。"""
    base = os.path.basename(filename).lower()
    if "512" in base:
        return "yolov3_512"
    if "416" in base:
        return "yolov3_416"
    return ""


def guess_input_from_filename(filename):
    """ファイル名から入力サイズを推定する。"""
    base = os.path.basename(filename).lower()
    m = re.search(r"(\d{3,4})x(\d{3,4})", base)
    if m:
        return f"{m.group(1)}x{m.group(2)}"
    if "512" in base:
        return "512x512"
    if "416" in base:
        return "416x416"
    return ""


# ---------- 評価 ----------

def eval_log_dual_iou(logfile, cfg, gt_dir, map_dir, workdir,
                      primary_iou, ref_iou):
    """1 本のログを主 IoU と参考 IoU で評価し、結果 dict を返す。"""
    n_frames = cfg["video"]["n_frames"]
    prefix = "frame_"
    digits = cfg["video"]["frame_pad"]
    mask_cfg = cfg["mask"]
    type_to_class = {int(k): v for k, v in cfg["classes"]["type_to_name"].items()}

    header, records = parse_log(logfile)

    # 座標スケール自動補正
    target_w = cfg["video"]["width"]
    target_h = cfg["video"]["height"]
    records, msg = auto_scale_if_needed(records, target_w, target_h)
    if msg:
        print(f"  {msg}")

    # 表示統計
    held_results, stats = build_held_results(records, n_frames)
    frame_results = build_frame_results(records, n_frames)

    # ヘッダから情報を取得（なければファイル名から推定 or 空欄）
    model = header.get("MODEL", "") or guess_model_from_filename(logfile)
    input_size = header.get("INPUT", "") or guess_input_from_filename(logfile)
    fps_str = header.get("FPS", "")
    try:
        fps = float(fps_str) if fps_str else None
    except ValueError:
        fps = None

    result = {
        "file": os.path.basename(logfile),
        "model": model,
        "input": input_size,
        "fps": fps,
        "show_every": header.get("SHOW_EVERY", "?"),
        "n_records": len(records),
        "n_valid": stats["n_valid"],
        "n_inversions": stats["n_inversions"],
        "avg_hold": stats["avg_hold"],
    }

    # detection-results を書き出し（1 回だけ）
    frame_det_dir = os.path.join(workdir, "det_frame")
    display_det_dir = os.path.join(workdir, "det_display")
    write_detection_results(
        frame_results, frame_det_dir, type_to_class, mask_cfg, prefix, digits,
    )
    write_detection_results(
        held_results, display_det_dir, type_to_class, mask_cfg, prefix, digits,
    )

    for iou in [primary_iou, ref_iou]:
        suffix = "75" if iou == 0.75 else "50"
        r_frame = run_map(map_dir, gt_dir, frame_det_dir, iou=iou)
        r_display = run_map(map_dir, gt_dir, display_det_dir, iou=iou)
        if r_frame:
            result[f"mAP_frame@{suffix}"] = r_frame["mAP"]
            for cls, ap in r_frame["per_class"].items():
                result[f"frame_{cls}@{suffix}"] = ap
        if r_display:
            result[f"mAP_display@{suffix}"] = r_display["mAP"]
            for cls, ap in r_display["per_class"].items():
                result[f"display_{cls}@{suffix}"] = ap
            # mAP(2cls): car + person のみ（bicycle を除く）
            ap_car = r_display["per_class"].get("car", 0)
            ap_person = r_display["per_class"].get("person", 0)
            result[f"mAP2_display@{suffix}"] = (ap_car + ap_person) / 2.0

    return result


# ---------- 画面出力 ----------

def print_report(results, primary_iou):
    """結果の一覧表を出力する。"""
    p = "75" if primary_iou == 0.75 else "50"
    r = "50" if p == "75" else "75"
    ref_iou = 0.5 if primary_iou != 0.5 else 0.75

    MODEL_W = 22
    # 列定義: (見出し, 幅)
    cols = [
        ("model",        MODEL_W, "left"),
        ("input",        8, "right"),
        ("SE",           3, "right"),
        ("FPS",          6, "right"),
        (f"car@.{p}",    7, "right"),
        (f"car@.{r}",    7, "right"),
        (f"person@.{p}", 10, "right"),
        (f"person@.{r}", 10, "right"),
        (f"bicycle@.{p}", 11, "right"),
        (f"bicycle@.{r}", 11, "right"),
        (f"mAP2@.{p}",   8, "right"),
        (f"mAP2@.{r}",   8, "right"),
        ("loss",         7, "right"),
        ("inv",          4, "right"),
    ]

    # 見出し行
    header_parts = []
    for name, w, align in cols:
        header_parts.append(pad_to_width(name, w, align))
    print("\n" + "  ".join(header_parts))
    total_w = sum(w for _, w, _ in cols) + 2 * (len(cols) - 1)
    print("-" * total_w)

    # データ行
    for res in results:
        model = shorten_model(res["model"] or res["file"], MODEL_W)
        inp = res["input"] or "---"
        fps = f"{res['fps']:.1f}" if res["fps"] is not None else "---"
        se = str(res["show_every"])

        car_p = res.get(f"display_car@{p}", 0)
        car_r = res.get(f"display_car@{r}", 0)
        per_p = res.get(f"display_person@{p}", 0)
        per_r = res.get(f"display_person@{r}", 0)
        bic_p = res.get(f"display_bicycle@{p}", 0)
        bic_r = res.get(f"display_bicycle@{r}", 0)
        m2_p = res.get(f"mAP2_display@{p}", 0)
        m2_r = res.get(f"mAP2_display@{r}", 0)

        # 損失 = mAP_display - mAP_frame（マイナスの値）
        mf = res.get(f"mAP_frame@{p}", 0)
        md = res.get(f"mAP_display@{p}", 0)
        loss = md - mf

        vals = [
            pad_to_width(model, MODEL_W, "left"),
            pad_to_width(inp, 8),
            pad_to_width(se, 3),
            pad_to_width(fps, 6),
            pad_to_width(f"{car_p:.1f}", 7),
            pad_to_width(f"{car_r:.1f}", 7),
            pad_to_width(f"{per_p:.1f}", 10),
            pad_to_width(f"{per_r:.1f}", 10),
            pad_to_width(f"{bic_p:.1f}", 11),
            pad_to_width(f"{bic_r:.1f}", 11),
            pad_to_width(f"{m2_p:.1f}", 8),
            pad_to_width(f"{m2_r:.1f}", 8),
            pad_to_width(f"{loss:+.1f}", 7),
            pad_to_width(f"{res['n_inversions']}", 4),
        ]
        print("  ".join(vals))

    print()
    print(f"  AP の値は mAP_display のクラス別 AP。")
    print(f"  @.{p} = IoU {primary_iou} (主指標)、@.{r} = IoU {ref_iou} (参考値)。")
    print(f"  mAP2 = (car + person) / 2。loss = mAP_display - mAP_frame (IoU {primary_iou})。")
    print(f"  SE = SHOW_EVERY（何枚に 1 枚表示したか）。inv = 表示順の逆行回数。")
    print()


# ---------- 列名の統一定義 ----------

def get_report_cols(primary_iou):
    """TSV / XLSX 共通の列名リストを返す。"""
    p = "75" if primary_iou == 0.75 else "50"
    r = "50" if p == "75" else "75"
    return [
        ("file", "file"),
        ("model", "model"),
        ("input", "input"),
        ("fps", "fps"),
        ("show_every", "SE"),
        ("n_records", "n_records"),
        (f"mAP_frame@{p}", f"mAP_frame@.{p}"),
        (f"mAP_display@{p}", f"mAP_display@.{p}"),
        (f"mAP2_display@{p}", f"mAP2@.{p}"),
        (f"frame_car@{p}", f"frame_car@.{p}"),
        (f"frame_person@{p}", f"frame_person@.{p}"),
        (f"frame_bicycle@{p}", f"frame_bicycle@.{p}"),
        (f"display_car@{p}", f"car@.{p}"),
        (f"display_person@{p}", f"person@.{p}"),
        (f"display_bicycle@{p}", f"bicycle@.{p}"),
        (f"mAP_frame@{r}", f"mAP_frame@.{r}"),
        (f"mAP_display@{r}", f"mAP_display@.{r}"),
        (f"mAP2_display@{r}", f"mAP2@.{r}"),
        (f"frame_car@{r}", f"frame_car@.{r}"),
        (f"frame_person@{r}", f"frame_person@.{r}"),
        (f"frame_bicycle@{r}", f"frame_bicycle@.{r}"),
        (f"display_car@{r}", f"car@.{r}"),
        (f"display_person@{r}", f"person@.{r}"),
        (f"display_bicycle@{r}", f"bicycle@.{r}"),
        ("n_valid", "n_valid"),
        ("n_inversions", "inv"),
        ("avg_hold", "avg_hold"),
    ]


# ---------- TSV / XLSX ----------

def write_tsv(results, path, primary_iou):
    """TSV ファイルに書き出す（BOM 付き UTF-8）。"""
    col_defs = get_report_cols(primary_iou)
    headers = [display for _, display in col_defs]
    keys = [key for key, _ in col_defs]

    with open(path, "w", encoding="utf-8-sig") as f:
        f.write("\t".join(headers) + "\n")
        for res in results:
            vals = []
            for k in keys:
                v = res.get(k, "")
                if isinstance(v, float):
                    vals.append(f"{v:.2f}")
                elif v is None:
                    vals.append("")
                else:
                    vals.append(str(v))
            f.write("\t".join(vals) + "\n")
    print(f"TSV 出力: {path}")


def write_xlsx(results, path, primary_iou):
    """Excel (.xlsx) ファイルに書き出す。"""
    try:
        from openpyxl import Workbook
    except ImportError:
        print("ERROR: openpyxl が必要です。pip install openpyxl を実行してください。",
              file=sys.stderr)
        return

    col_defs = get_report_cols(primary_iou)
    headers = [display for _, display in col_defs]
    keys = [key for key, _ in col_defs]

    wb = Workbook()
    ws = wb.active
    ws.title = "report"

    # 見出し行
    for c_idx, h in enumerate(headers, 1):
        ws.cell(row=1, column=c_idx, value=h)

    # データ行
    for r_idx, res in enumerate(results, 2):
        for c_idx, k in enumerate(keys, 1):
            v = res.get(k, "")
            if v is None:
                v = ""
            ws.cell(row=r_idx, column=c_idx, value=v)

    # 見出し行を固定
    ws.freeze_panes = "A2"

    wb.save(path)
    print(f"Excel 出力: {path}")


# ---------- メイン ----------

def main():
    parser = argparse.ArgumentParser(
        description="複数の学生のログをまとめて一覧にする"
    )
    parser.add_argument("logfiles", nargs="+", help="detection_log.txt のパス（複数可）")
    parser.add_argument(
        "--config",
        default=os.path.join(os.path.dirname(__file__), "..", "config.yaml"),
        help="config.yaml のパス",
    )
    parser.add_argument("--gt-dir", default=None, help="GT ディレクトリ")
    parser.add_argument("--tsv", default=None, help="TSV ファイルにも出力する場合のパス (BOM 付き UTF-8)")
    parser.add_argument("--xlsx", default=None, help="Excel ファイルにも出力する場合のパス")
    args = parser.parse_args()

    cfg = load_config(args.config)
    base = os.path.dirname(os.path.abspath(args.config))
    eval_dir = os.path.dirname(os.path.abspath(__file__))

    primary_iou = cfg.get("eval", {}).get("iou_threshold", 0.75)
    ref_iou = 0.5 if primary_iou != 0.5 else 0.75
    gt_dir = args.gt_dir or os.path.join(base, cfg["paths"]["ground_truth"])
    map_dir = os.path.join(base, cfg["paths"]["map_tool"])
    workdir = os.path.join(eval_dir, "work", "report")

    results = []
    for i, logfile in enumerate(sorted(args.logfiles)):
        print(f"[{i+1}/{len(args.logfiles)}] {logfile} ...", flush=True)
        res = eval_log_dual_iou(
            logfile, cfg, gt_dir, map_dir, workdir,
            primary_iou, ref_iou,
        )
        results.append(res)

    print_report(results, primary_iou)

    if args.tsv:
        write_tsv(results, args.tsv, primary_iou)

    if args.xlsx:
        write_xlsx(results, args.xlsx, primary_iou)


if __name__ == "__main__":
    main()
